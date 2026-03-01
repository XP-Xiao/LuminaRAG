"""Excel(.xlsx) 结构化提取：按 sheet 拆成多「页」供三层分块复用。"""
from __future__ import annotations

from typing import Any

import openpyxl
from langchain_core.documents import Document


def parse_xlsx_to_sheets(file_path: str) -> list[dict[str, Any]]:
    """按 sheet 拆页；空 sheet 跳过。返回 [{"page", "title", "text"}]。"""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        sections: list[dict[str, Any]] = []
        for idx, ws in enumerate(wb.worksheets, start=1):
            rows_text: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    rows_text.append(" ".join(cells))
            text = "\n".join(rows_text).strip()
            if not text:
                continue
            sections.append({"page": idx, "title": ws.title, "text": text})
        return sections
    finally:
        wb.close()


def load_xlsx_for_document_loader(file_path: str, filename: str) -> list[Document]:
    """供 DocumentLoader 使用：返回 LangChain Document 列表，metadata.page 为 sheet 序号。"""
    docs: list[Document] = []
    for sec in parse_xlsx_to_sheets(file_path):
        docs.append(
            Document(
                page_content=sec["text"],
                metadata={
                    "page": sec["page"],
                    "source": filename,
                    "section_title": sec["title"],
                },
            )
        )
    return docs
